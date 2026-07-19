from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.exports import gerar_excel, gerar_pdf


def test_excel_contem_dados_memoria_e_metadados():
    dados = pd.DataFrame([{"Distribuidora": "LIGHT", "TUSD": 100.0}])
    memoria = pd.DataFrame([{"Componente": "TE", "Subtotal": 200.0}])
    conteudo = gerar_excel(dados, memoria, "LIGHT | A4")
    workbook = load_workbook(BytesIO(conteudo), read_only=True)
    assert workbook.sheetnames == ["Dados ANEEL", "Memória de cálculo", "Metadados"]


def test_pdf_tem_assinatura_valida():
    memoria = pd.DataFrame([{"Componente": "TE", "Subtotal": 200.0}])
    assert gerar_pdf(memoria, "LIGHT | A4", 200.0).startswith(b"%PDF")
